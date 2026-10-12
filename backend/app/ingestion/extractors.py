"""Text extraction for supported document types (SRS FR-006/FR-008).

Each extractor returns a list of :class:`ExtractedPage` so page-level metadata
can flow through to citations. DOCX and TXT files, which have no reliable page
boundaries, are represented as a single page.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from app.core.exceptions import IngestionError


@dataclass(slots=True)
class ExtractedPage:
    page_number: int | None
    text: str


@dataclass(slots=True)
class ExtractedDocument:
    pages: list[ExtractedPage]
    page_count: int | None

    @property
    def text(self) -> str:
        return "\n\n".join(page.text for page in self.pages if page.text.strip())


def _extract_pdf(path: Path) -> ExtractedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
    except (PdfReadError, OSError) as exc:
        raise IngestionError("The PDF file could not be parsed.", code="pdf_parse_error") from exc

    pages: list[ExtractedPage] = []
    for index, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - individual page failures shouldn't abort ingestion
            text = ""
        pages.append(ExtractedPage(page_number=index + 1, text=text))

    if not any(page.text.strip() for page in pages):
        raise IngestionError(
            "No extractable text found. The PDF may be scanned; OCR is not enabled.",
            code="pdf_no_text",
        )
    return ExtractedDocument(pages=pages, page_count=len(pages))


def _extract_docx(path: Path) -> ExtractedDocument:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:  # pragma: no cover
        raise IngestionError("DOCX support is unavailable.", code="docx_unavailable") from exc

    try:
        document = DocxDocument(str(path))
    except Exception as exc:  # noqa: BLE001 - python-docx raises varied errors
        raise IngestionError("The DOCX file could not be parsed.", code="docx_parse_error") from exc

    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    text = "\n".join(part for part in parts if part.strip())
    if not text.strip():
        raise IngestionError("The DOCX file contains no extractable text.", code="docx_no_text")
    return ExtractedDocument(pages=[ExtractedPage(page_number=None, text=text)], page_count=None)


def _extract_txt(path: Path) -> ExtractedDocument:
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        text = path.read_text(encoding="latin-1")
    except OSError as exc:
        raise IngestionError("The text file could not be read.", code="txt_read_error") from exc
    if not text.strip():
        raise IngestionError("The text file is empty.", code="txt_empty")
    return ExtractedDocument(pages=[ExtractedPage(page_number=None, text=text)], page_count=None)


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".txt": _extract_txt,
}


def extract_document(path: Path, file_type: str) -> ExtractedDocument:
    """Extract text from a file based on its normalized extension (e.g. ``.pdf``)."""
    extractor = _EXTRACTORS.get(file_type.lower())
    if extractor is None:
        raise IngestionError(
            f"Unsupported file type: {file_type}", code="unsupported_file_type"
        )
    return extractor(path)
